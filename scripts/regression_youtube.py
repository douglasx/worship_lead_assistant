from __future__ import annotations

import argparse
import sys
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.main import app


def get_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run YouTube -> PPTX regression test.")
    parser.add_argument(
        "--url",
        default="https://youtu.be/n0FBb6hnwTo?si=vlEZ8asfdAp1S9oW",
        help="YouTube URL.",
    )
    parser.add_argument(
        "--source",
        default="all",
        choices=["transcript", "description", "auto", "all"],
        help="Source mode to test. 'all' runs transcript, description, and auto.",
    )
    parser.add_argument("--lines", type=int, default=4, help="Lines per slide.")
    parser.add_argument("--font-size", type=int, default=56, help="Lyric font size.")
    parser.add_argument(
        "--vocal-part",
        default="everyone",
        choices=["male", "female", "everyone"],
        help="Title/counter vocal-part color mode.",
    )
    parser.add_argument(
        "--footer-note",
        default="無法久站者請自行坐下",
        help="Footer note for first slide.",
    )
    parser.add_argument(
        "--output-prefix",
        default="samples/regression_youtube_output",
        help="Output PPTX path prefix (source suffix will be added).",
    )
    return parser.parse_args()


def inspect_ppt(content: bytes) -> dict:
    prs = Presentation(BytesIO(content))
    result = {
        "slide_count": len(prs.slides),
        "background_rgb": "N/A",
        "first_text": "",
    }
    if not prs.slides:
        return result

    slide = prs.slides[0]
    try:
        fill = slide.background.fill
        if fill is not None and getattr(fill, "type", None) is not None:
            color = fill.fore_color
            rgb = getattr(color, "rgb", None)
            if rgb is not None:
                result["background_rgb"] = str(rgb)
    except Exception:
        pass

    for shp in slide.shapes:
        if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
            txt = (shp.text or "").strip()
            if txt:
                result["first_text"] = txt.replace("\n", " ")[:200]
                break

    return result


def run_source(client: TestClient, args: argparse.Namespace, source: str) -> tuple[bool, Path | None]:
    data = {
        "urls": args.url,
        "source": source,
        "lines_per_slide": str(args.lines),
        "font_size": str(args.font_size),
        "vocal_part": args.vocal_part,
        "footer_note": args.footer_note,
    }
    resp = client.post("/api/convert-youtube", data=data)

    print(f"KEY source={source} status_code={resp.status_code} output_bytes={len(resp.content)}")

    if resp.status_code != 200:
        text = resp.text
        if len(text) > 400:
            text = text[:400] + "..."
        print(f"KEY source={source} error={text}")
        return False, None

    out = Path(f"{args.output_prefix}_{source}.pptx")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_bytes(resp.content)
    inspect = inspect_ppt(resp.content)

    print(f"KEY source={source} saved_file={out}")
    print(f"KEY source={source} slide_count={inspect['slide_count']}")
    print(f"KEY source={source} first_slide_background_rgb={inspect['background_rgb']}")
    print(f"KEY source={source} first_text={inspect['first_text']}")
    return True, out


def main() -> int:
    args = get_args()
    client = TestClient(app)

    modes = [args.source] if args.source != "all" else ["transcript", "description", "auto"]
    success = False

    for mode in modes:
        ok, _ = run_source(client, args, mode)
        success = success or ok

    if success:
        print("PASS/FAIL: PASS")
        return 0

    print("PASS/FAIL: FAIL")
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
