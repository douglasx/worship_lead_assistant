from __future__ import annotations

import argparse
import sys
from io import BytesIO
from pathlib import Path

from fastapi.testclient import TestClient
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.main import app


def get_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Run PDF -> PPTX regression test.")
    parser.add_argument(
        "--input",
        default="samples/有一位神（D）.pdf",
        help="Input PDF path.",
    )
    parser.add_argument(
        "--output",
        default="samples/regression_pdf_output.pptx",
        help="Output PPTX path.",
    )
    parser.add_argument("--lines", type=int, default=4, help="Lines per slide.")
    parser.add_argument("--font-size", type=int, default=56, help="Lyric font size.")
    parser.add_argument(
        "--vocal-part",
        default="everyone",
        choices=["male", "female", "everyone"],
        help="Title/counter vocal-part color mode.",
    )
    parser.add_argument(
        "--footer-note",
        default="無法久站者請自行坐下",
        help="Footer note for first slide.",
    )
    return parser.parse_args()


def inspect_ppt(content: bytes) -> dict:
    prs = Presentation(BytesIO(content))
    result = {
        "slide_count": len(prs.slides),
        "background_rgb": "N/A",
        "first_text": "",
    }
    if not prs.slides:
        return result

    slide = prs.slides[0]
    try:
        fill = slide.background.fill
        if fill is not None and getattr(fill, "type", None) is not None:
            color = fill.fore_color
            rgb = getattr(color, "rgb", None)
            if rgb is not None:
                result["background_rgb"] = str(rgb)
    except Exception:
        pass

    for shp in slide.shapes:
        if getattr(shp, "has_text_frame", False) and shp.has_text_frame:
            txt = (shp.text or "").strip()
            if txt:
                result["first_text"] = txt.replace("\n", " ")[:200]
                break

    return result


def main() -> int:
    args = get_args()
    input_path = Path(args.input)
    output_path = Path(args.output)

    if not input_path.exists():
        print(f"KEY status=FAIL reason=input_missing path={input_path}")
        return 1

    output_path.parent.mkdir(parents=True, exist_ok=True)

    client = TestClient(app)
    with input_path.open("rb") as f:
        files = [("files", (input_path.name, f, "application/pdf"))]
        data = {
            "lines_per_slide": str(args.lines),
            "font_size": str(args.font_size),
            "vocal_part": args.vocal_part,
            "footer_note": args.footer_note,
        }
        resp = client.post("/api/convert", files=files, data=data)

    print(f"KEY status_code={resp.status_code}")
    print(f"KEY content_type={resp.headers.get('content-type')}")
    print(f"KEY content_disposition={resp.headers.get('content-disposition')}")
    print(f"KEY output_bytes={len(resp.content)}")

    if resp.status_code != 200:
        text = resp.text
        if len(text) > 400:
            text = text[:400] + "..."
        print(f"KEY error={text}")
        print("PASS/FAIL: FAIL")
        return 1

    output_path.write_bytes(resp.content)
    inspect = inspect_ppt(resp.content)

    print(f"KEY saved_file={output_path}")
    print(f"KEY slide_count={inspect['slide_count']}")
    print(f"KEY first_slide_background_rgb={inspect['background_rgb']}")
    print(f"KEY first_text={inspect['first_text']}")
    print("PASS/FAIL: PASS")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
